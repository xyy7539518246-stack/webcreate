# -*- coding: utf-8 -*-
import docx
doc = docx.Document(r'D:\webcreat\webcreate\docs\WebCreate项目实践报告_v2.docx')

print('=== 封面段落 ===')
for i in range(12):
    p = doc.paragraphs[i]
    text = p.text.strip()
    if not text:
        continue
    runs_info = ''
    if p.runs:
        r = p.runs[0]
        sz = r.font.size.pt if r.font.size else '?'
        nm = r.font.name
        bd = r.font.bold
        runs_info = f'[{nm} {sz}pt bold={bd}]'
    print(f'[{i}] align={p.alignment} {runs_info}')
    print(f'     {text[:60]}')

print('\n=== 封面信息表格（表格0）===')
t = doc.tables[0]
print(f'行数: {len(t.rows)}, 列数: {len(t.columns)}')
for row_idx, row in enumerate(t.rows):
    cells_text = [c.text.strip() for c in row.cells]
    print(f'  行{row_idx}: {" | ".join(cells_text)}')

print(f'\n总表格数: {len(doc.tables)}')
import os
f = os.path.getsize(r'D:\webcreat\webcreate\docs\WebCreate项目实践报告_v2.docx')
print(f'文件大小: {f/1024:.1f} KB')
