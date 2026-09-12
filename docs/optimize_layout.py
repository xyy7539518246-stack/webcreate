# -*- coding: utf-8 -*-
"""WebCreate 项目实践报告排版优化"""
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INPUT = r'D:\webcreat\webcreate\docs\WebCreate项目实践报告_v2.docx'
OUTPUT = r'D:\webcreat\webcreate\docs\WebCreate项目实践报告_v2.docx'

doc = Document(INPUT)

# ========== 工具函数 ==========
def set_run_font(run, name='宋体', size=12, bold=False, color=None):
    run.font.name = name
    run.element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color

def set_cell_shading(cell, color_hex):
    """设置单元格底色"""
    tcPr = cell._tc.get_or_add_tcPr()
    # 移除已有的 shd
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def set_table_width_pct(table, pct=100):
    """设置表格宽度百分比"""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    for old in tblPr.findall(qn('w:tblW')):
        tblPr.remove(old)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(pct * 50)))
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)

# ========== 1. 页面设置：A4 + 标准页边距 ==========
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

# ========== 2. 基础样式定义 ==========
# Normal 正文
normal = doc.styles['Normal']
normal.font.name = '宋体'
normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)

# 标题样式
heading_configs = {
    1: {'size': 16, 'before': 24, 'after': 18, 'align': WD_ALIGN_PARAGRAPH.CENTER},
    2: {'size': 14, 'before': 12, 'after': 6, 'align': WD_ALIGN_PARAGRAPH.LEFT},
    3: {'size': 12, 'before': 6, 'after': 3, 'align': WD_ALIGN_PARAGRAPH.LEFT},
}
for level, cfg in heading_configs.items():
    hs = doc.styles[f'Heading {level}']
    hs.font.name = '黑体'
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    hs.font.size = Pt(cfg['size'])
    hs.font.bold = True
    hs.font.color.rgb = RGBColor(0, 0, 0)
    hs.paragraph_format.space_before = Pt(cfg['before'])
    hs.paragraph_format.space_after = Pt(cfg['after'])
    hs.paragraph_format.alignment = cfg['align']
    hs.paragraph_format.first_line_indent = Pt(0)
    hs.paragraph_format.line_spacing = 1.5

# ========== 3. 区域识别与段落精细排版 ==========
# 先确定各区域起始索引
region_markers = {}
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    s = p.style.name
    # 只匹配 Heading 1 级别的标题，避免目录中的同名文本干扰
    if s == 'Heading 1':
        if t == '摘　　要':
            region_markers['abstract'] = i
        elif t == '目　　录':
            region_markers['toc'] = i
        elif t.startswith('1、'):
            region_markers['body'] = i
        elif t == '参考文献':
            region_markers['refs'] = i
        elif t == '附录':
            region_markers['appendix'] = i

print(f'区域标记: {region_markers}')

def get_region(idx):
    if idx < region_markers.get('abstract', 999):
        return 'cover'
    elif idx < region_markers.get('toc', 999):
        return 'abstract'
    elif idx < region_markers.get('body', 999):
        return 'toc'
    elif idx < region_markers.get('refs', 999):
        return 'body'
    elif idx < region_markers.get('appendix', 999):
        return 'refs'
    else:
        return 'appendix'

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    style = p.style.name
    region = get_region(i)
    pf = p.paragraph_format

    # ---- 封面区域（校名 + 实习实践报告 + 信息表格） ----
    if region == 'cover':
        pf.first_line_indent = Pt(0)
        if text == '河北科技大学':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = Pt(24)
            pf.space_after = Pt(36)
            for r in p.runs:
                r.font.name = '楷体'
                r.element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')
                r.font.size = Pt(28)
                r.font.bold = True
        elif text == '实习实践报告':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_after = Pt(48)
            for r in p.runs:
                set_run_font(r, '黑体', 42, True)
        elif not text:
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
        continue

    # ---- 摘要区域 ----
    if region == 'abstract':
        if style.startswith('Heading'):
            # 摘要标题：参考样例 黑体18pt加粗居中
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = Pt(12)
            pf.space_after = Pt(12)
            pf.first_line_indent = Pt(0)
            for r in p.runs:
                set_run_font(r, '黑体', 18, True)
            continue
        if text.startswith('关键词'):
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(12)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                if '关键词' in r.text:
                    r.font.bold = True
            continue
        if text:
            pf.first_line_indent = Pt(24)
            pf.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        continue

    # ---- 目录区域 ----
    if region == 'toc':
        if style.startswith('Heading'):
            # 目录标题：参考样例 黑体18pt加粗居中
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = Pt(12)
            pf.space_after = Pt(12)
            pf.first_line_indent = Pt(0)
            for r in p.runs:
                set_run_font(r, '黑体', 18, True)
            continue
        if text:
            pf.first_line_indent = Pt(0)
            pf.left_indent = Cm(2)
            pf.space_after = Pt(6)
            pf.line_spacing = 1.5
            for r in p.runs:
                set_run_font(r, '宋体', 14, False)
        continue

    # ---- 正文区域 ----
    if region == 'body':
        if style.startswith('Heading'):
            pf.first_line_indent = Pt(0)
            continue
        if not text:
            continue
        # 普通正文：首行缩进2字符，1.5倍行距，两端对齐
        pf.first_line_indent = Pt(24)
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            if r.font.name is None or r.font.name == '':
                set_run_font(r, '宋体', 12, False)
            elif r.font.size is None:
                r.font.size = Pt(12)
        continue

    # ---- 参考文献区域 ----
    if region == 'refs':
        if style.startswith('Heading'):
            continue
        if text.startswith('[') and ']' in text[:6]:
            pf.first_line_indent = Pt(0)
            pf.left_indent = Cm(0.74)
            pf.hanging_indent = Cm(0.74)
            pf.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for r in p.runs:
                set_run_font(r, '宋体', 10.5, False)
        continue

    # ---- 附录区域 ----
    if region == 'appendix':
        if style.startswith('Heading'):
            continue
        # 判断是否为目录结构代码块
        is_code = (
            '├──' in text or '└──' in text or
            text.strip().startswith('#') or
            'webcreate/' in text or
            (text and text[0] in '│├└' )
        )
        if is_code and text:
            pf.first_line_indent = Pt(0)
            pf.line_spacing = 1.0
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            for r in p.runs:
                r.font.name = 'Consolas'
                r.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                r.font.size = Pt(9)
        elif text:
            pf.first_line_indent = Pt(24)
            pf.line_spacing = 1.5
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        continue

# ========== 4. 表格排版优化 ==========
HEADER_BG = 'D9E2F3'  # 浅蓝表头
ALT_BG = 'F7F9FC'     # 极浅蓝交替行

for ti, table in enumerate(doc.tables):
    # 第一个表格是封面信息表（表单样式），跳过数据表优化
    if ti == 0:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        continue

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width_pct(table, 100)

    n_rows = len(table.rows)
    n_cols = len(table.columns)

    for row_idx, row in enumerate(table.rows):
        # 设置行高
        trPr = row._tr.get_or_add_trPr()
        # 移除旧的 trHeight
        for old in trPr.findall(qn('w:trHeight')):
            trPr.remove(old)
        trHeight = OxmlElement('w:trHeight')
        trHeight.set(qn('w:val'), '400')  # 最小行高 20pt
        trHeight.set(qn('w:hRule'), 'atLeast')
        trPr.append(trHeight)

        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.first_line_indent = Pt(0)
                pf.line_spacing = 1.3
                pf.space_before = Pt(3)
                pf.space_after = Pt(3)

                for r in p.runs:
                    set_run_font(r, '宋体', 10.5, False)

            if row_idx == 0:
                # 表头
                set_cell_shading(cell, HEADER_BG)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.bold = True
                        set_run_font(r, '黑体', 10.5, True)
            else:
                # 数据行
                # 交替行底色（大表格启用，小表格不用）
                if n_rows > 8 and row_idx % 2 == 0:
                    set_cell_shading(cell, ALT_BG)
                # 第一列居中，其余左对齐
                for p in cell.paragraphs:
                    if col_idx == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # 特殊：git 提交记录表（33行），时间列也居中
    if n_rows > 20:
        for row_idx in range(1, n_rows):
            for col_idx in [0, 1]:  # 哈希和时间列居中
                for p in table.rows[row_idx].cells[col_idx].paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ========== 5. 保存 ==========
doc.save(OUTPUT)
print(f'排版优化完成: {OUTPUT}')
print(f'总段落数: {len(doc.paragraphs)}')
print(f'表格数: {len(doc.tables)}')
