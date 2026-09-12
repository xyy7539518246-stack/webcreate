# -*- coding: utf-8 -*-
"""生成 WebCreate 项目实践报告 docx"""
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ========== 全局样式设置 ==========
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.first_line_indent = Cm(0.74)

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = '黑体'
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    hs.font.color.rgb = RGBColor(0, 0, 0)
    if level == 1:
        hs.font.size = Pt(16)
    elif level == 2:
        hs.font.size = Pt(14)
    else:
        hs.font.size = Pt(13)

def add_para(text, bold=False, align=None, indent=True, size=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    if bold:
        run.bold = True
    if size:
        run.font.size = Pt(size)
    if align:
        p.alignment = align
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    return p

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.first_line_indent = Cm(0)
    return h

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 表头
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run.font.size = Pt(10.5)
    # 数据行
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.name = '宋体'
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.size = Pt(10.5)
    return table

# ========== 封面（参考样例：校名 + 实习实践报告 + 信息表格） ==========
# 顶部校名（用户可自行替换为校徽+校名图片）
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Cm(0)
p.paragraph_format.space_before = Pt(24)
p.paragraph_format.space_after = Pt(36)
run = p.add_run('河北科技大学')
run.font.name = '楷体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')
run.font.size = Pt(28)
run.bold = True

# 主标题
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Cm(0)
p.paragraph_format.space_after = Pt(48)
run = p.add_run('实习实践报告')
run.font.name = '黑体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
run.font.size = Pt(42)
run.bold = True

# 信息表格（带下划线的表单样式）
info_data = [
    ('学生姓名：', '胥阳', '学　　号：', '________'),
    ('专业班级：', '计算机科学与技术 2505 班', '', ''),
    ('实习类别：', '☑认识实习　□生产实习　□毕业实习', '', ''),
    ('实习方式：', '☑集中实习　□分散实习', '', ''),
    ('实习单位：', '河北科技大学', '', ''),
    ('指导教师：', '黄敏', '', ''),
    ('实习时间：', '第 1 周到第 1 周，共 1 周', '', '9月4日—9月6日'),
    ('所属学院：', '信息科学与工程学院', '', ''),
]

# 用表格实现，4列：标签1 | 内容1 | 标签2 | 内容2
table = doc.add_table(rows=len(info_data), cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
# 设置表格宽度
tbl = table._tbl
tblPr = tbl.tblPr
tblW = OxmlElement('w:tblW')
tblW.set(qn('w:w'), '9000')
tblW.set(qn('w:type'), 'dxa')
for old in tblPr.findall(qn('w:tblW')):
    tblPr.remove(old)
tblPr.append(tblW)

for row_idx, (label1, value1, label2, value2) in enumerate(info_data):
    row = table.rows[row_idx]
    # 设置行高
    trPr = row._tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), '500')
    trHeight.set(qn('w:hRule'), 'atLeast')
    trPr.append(trHeight)

    cells = [
        (row.cells[0], label1, True),   # 标签1 右对齐
        (row.cells[1], value1, False),  # 内容1 左对齐 下划线
        (row.cells[2], label2, True),   # 标签2 右对齐
        (row.cells[3], value2, False),  # 内容2 左对齐 下划线
    ]
    for cell, text, is_label in cells:
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        if is_label:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(text)
            run.font.name = '宋体'
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.size = Pt(14)
            run.bold = True
            # 标签单元格无边框
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            for border_name in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement(f'w:{border_name}')
                border.set(qn('w:val'), 'nil')
                tcBorders.append(border)
            tcPr.append(tcBorders)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = '宋体'
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.size = Pt(14)
            # 内容单元格只有底部边框（下划线效果）
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            for border_name in ['top', 'left', 'right']:
                border = OxmlElement(f'w:{border_name}')
                border.set(qn('w:val'), 'nil')
                tcBorders.append(border)
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '0')
            bottom.set(qn('w:color'), '000000')
            tcBorders.append(bottom)
            tcPr.append(tcBorders)

# 设置列宽
col_widths = [Cm(2.8), Cm(5.5), Cm(2.5), Cm(4.0)]
for row in table.rows:
    for idx, width in enumerate(col_widths):
        row.cells[idx].width = width

doc.add_page_break()

# ========== 摘要 ==========
add_heading('摘　　要', level=1)

add_para('本实践报告围绕"WebCreate —— AI 编程学习助手"前端单页应用的完整开发过程展开。本项目基于 Vue 3（Composition API）+ Vite 5 + Pinia + Vue Router 4 技术栈，采用纯前端实现方案，全部业务数据依靠本地 JSON 资源文件与浏览器 localStorage 完成持久化，实现了用户登录注册、编程课程学习、在线题库练习、代码示例查阅、AI 编程问答助手、个人学习中心等完整业务模块。')

add_para('项目使用 Git 进行版本控制，采用 develop—test—main 三分支开发流程，借助 GitHub Actions CI 流水线完成项目自动构建，并部署托管于 GitHub Pages 静态平台。开发过程中充分考虑跨浏览器兼容、前端性能优化、演示数据合规伦理、AI 工具使用边界等复杂工程约束；完成了需求分析、架构设计、编码实现、问题调试、线上部署、功能测试的全流程。')

add_para('通过本次实践，系统掌握了 Vue 3 组件化开发、SPA 单页应用路由处理与守卫、Pinia 状态管理、静态站点部署与 CI 自动化等核心 Web 开发技能，同时深入理解了纯前端静态项目的固有局限，明确了后续迭代优化方向。')

p = doc.add_paragraph()
p.paragraph_format.first_line_indent = Cm(0)
run = p.add_run('关键词：')
run.bold = True
run.font.name = '宋体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
run = p.add_run('Vue 3；Vite；SPA；Pinia；GitHub Pages；前端静态应用；AI 编程助手')
run.font.name = '宋体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

doc.add_page_break()

# ========== 目录 ==========
add_heading('目　　录', level=1)
toc_items = [
    '1、需求与选题分析',
    '2、可行性分析（含复杂工程约束考量）',
    '3、系统架构设计',
    '4、环境搭建与工具使用',
    '5、开发过程与问题解决',
    '6、部署与功能测试',
    '7、工具/框架局限性分析',
    '8、总结与改进方向',
    '参考文献',
    '附录',
]
for item in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(item)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(12)

doc.add_page_break()

# ========== 第1章 ==========
add_heading('1、需求与选题分析', level=1)

add_heading('1.1 项目背景', level=2)
add_para('在 Web 开发课程学习过程中，编程初学者往往面临学习资源分散、练习题缺乏即时反馈、代码样例不易查阅、遇到问题难以快速获得解答等痛点。现有学习平台通常需要后端服务器支持，部署和维护成本较高，不适合课程实践中快速上线验证。')
add_para('基于上述背景，本项目选题为"WebCreate —— AI 编程学习助手"，面向高校计算机专业学生、编程初学者与竞赛备赛者，提供"学习资源 → 题库练习 → 示例代码 → AI 问答"一体化的学习路径。项目以 Web 前端为核心，不需要后端服务器，浏览器即可直接运行，实现一站式编程学习体验。')

add_heading('1.2 功能需求', level=2)
add_para('根据选题定位，系统需实现以下功能模块：')
add_para('（1）用户认证模块：支持用户注册、登录、忘记密码（验证码重置），表单校验（手机号格式、验证码、密码强度），登录状态路由守卫，会话保存在 localStorage，支持退出登录。', indent=True)
add_para('（2）首页模块：系统欢迎首页，支持选择学习方向并展示对应学习路线，提供各功能模块导航入口。', indent=True)
add_para('（3）学习资源模块：加载本地 courses.json 课程数据，提供 Web 知识点分类导航、学习路线图时间线展示、课程与竞赛内容浏览，支持按资源类型动态过滤。', indent=True)
add_para('（4）题库练习模块：读取 quiz.json 题库资源，支持按知识点抽题练习、即时判题、答题记录、题目收藏、错题收集保存到本地存储、错题本查看、抽题循环玩法。', indent=True)
add_para('（5）示例代码模块：读取 examples.json，展示 C++ / Java / 前端多语言示例代码，支持分类切换、语法高亮、一键复制。', indent=True)
add_para('（6）AI 问答助手模块：内置本地知识库 assistant.json（45 条知识条目），优先本地知识库检索回答；可对接 DeepSeek 大模型 API 获取 AI 回答；API 异常或未配置 Key 时自动降级使用本地知识库；API 密钥通过页面配置保存在浏览器 localStorage，不硬编码到源码。', indent=True)
add_para('（7）个人中心模块：查看个人信息、学习方向、学习收藏、错题记录、学习记录，管理本地存储的用户学习数据。', indent=True)

add_heading('1.3 非功能需求', level=2)
add_para('（1）浏览器兼容与响应式：页面在 Chrome、Edge 等主流浏览器下布局与交互正常，完成全站移动端响应式适配，支持手机与电脑端访问。', indent=True)
add_para('（2）性能：JSON 本地数据加载流畅，Vite 构建完成代码分割（vue 核心库独立 chunk）与资源压缩；页面组件按需懒加载，首屏加载无明显卡顿。', indent=True)
add_para('（3）部署：SPA 单页应用部署 GitHub Pages，解决深层路由刷新 404 问题，线上环境可稳定访问。', indent=True)
add_para('（4）数据安全：全部为演示教学数据，不采集用户隐私信息；API 密钥存储于浏览器本地，不会提交至代码仓库或上传服务器。', indent=True)
add_para('（5）工程约束：单人独立开发，Git 完整版本记录，三分支管理，CI 自动化构建部署。', indent=True)

add_heading('1.4 风险识别', level=2)
add_para('（1）SPA 部署 GitHub Pages 路由刷新 404：静态托管不原生支持前端路由 history 模式，深层路由刷新会返回 404。', indent=True)
add_para('（2）大量 JSON 静态数据加载与页面渲染：题库和知识库数据量较大，一次性渲染可能导致页面卡顿。', indent=True)
add_para('（3）AI 接口网络异常或 Key 未配置：DeepSeek API 调用可能因网络、额度或 Key 问题失败，程序需要降级兜底。', indent=True)
add_para('（4）登录路由守卫处理不当：未登录用户可能通过直接输入 URL 访问受保护页面。', indent=True)
add_para('（5）本地存储数据一致性：多页面共享 localStorage 数据，需处理数据读写的时序与持久化。', indent=True)

doc.add_page_break()

# ========== 第2章 ==========
add_heading('2、可行性分析（含复杂工程约束考量）', level=1)

add_heading('2.1 总体技术方案', level=2)
add_para('技术栈：HTML5 + CSS3 + JavaScript（ES6+）+ Vue 3（Composition API）+ Vite 5 + Vue Router 4 + Pinia；')
add_para('开发工具：VS Code（Volar / ESLint / Prettier）；版本管理：Git + GitHub；部署：GitHub Actions + GitHub Pages；')
add_para('数据方案：本地 JSON 文件（courses.json / quiz.json / examples.json / assistant.json）提供业务数据，localStorage 完成用户状态、错题、收藏、学习方向、API 密钥的持久化；无后端服务。')

add_heading('2.2 选型理由', level=2)
add_para('（1）Vue 3 组合式 API（Composition API）便于组件拆分与逻辑复用，响应式系统成熟，适合快速开发单页应用；Vite 构建工具基于原生 ESM，热更新速度快，构建产物优化良好，适合短期课程实践项目。', indent=True)
add_para('（2）Pinia 实现全局状态管理，作为 Vue 官方推荐的状态管理方案，API 简洁，TypeScript 支持良好，替代 Vuex 简化了登录状态、用户数据的全局管理。', indent=True)
add_para('（3）Vue Router 实现 SPA 路由，配合路由守卫（beforeEach）实现页面权限控制，支持 history 模式与滚动行为定制。', indent=True)
add_para('（4）GitHub Pages 提供免费静态托管，不需要购买服务器；GitHub Actions 提供免费 CI 流水线，代码推送自动构建上线，三分支工作流实现开发—测试—发布的规范化管理。', indent=True)
add_para('（5）不开发后端服务，全部前端实现，降低项目复杂度与部署成本，适合 1 周课程实践完成；本地 JSON + localStorage 方案足以支撑教学演示级别的数据需求。', indent=True)

add_heading('2.3 功能模块划分', level=2)
add_table(
    ['模块', '对应页面/路由', '主要功能'],
    [
        ['登录认证', '/login, /register, /forgot-password', '手机号注册、登录、验证码重置密码、表单校验、路由守卫、会话保存、退出登录'],
        ['首页', '/', '系统欢迎页、学习方向选择、功能导航入口'],
        ['学习资源', '/courses', '知识点分类、学习路线时间线、课程/竞赛浏览、类型动态过滤'],
        ['题库练习', '/quiz', '按知识点抽题、即时判题、错题收集、题目收藏、抽题循环'],
        ['示例代码', '/examples', 'C++/Java/前端示例、分类切换、语法高亮、一键复制'],
        ['AI 问答', '/assistant', '本地知识库检索（45条）、DeepSeek API 直连、Key 页面配置、降级提示、对话记录'],
        ['个人中心', '/profile', '个人信息、学习方向、收藏、错题本、学习记录'],
    ]
)

add_heading('2.4 复杂工程约束可行性分析', level=2)

add_heading('2.4.1 跨浏览器兼容与响应式约束', level=3)
add_para('项目使用 Flex、Grid 标准 CSS 布局，配合 CSS reset 样式，避免浏览器私有特性；Vue 标准组件写法与 ES6+ 语法在现代浏览器中均有良好支持。项目在开发后期完成了全站响应式适配（git 提交 54544a6），通过媒体查询（@media max-width: 768px）调整内边距、布局方向与字号，确保移动端屏幕自动适配。该约束技术可行，已在 Chrome、Edge 浏览器实测通过。')

add_heading('2.4.2 性能约束', level=3)
add_para('项目业务数据使用 JSON 静态文件，通过 Vite 的 import 机制在组件中按需加载；Vite 生产构建自动开启 esbuild 压缩与代码分割（manualChunks 将 vue、vue-router、pinia 独立为 vendor chunk）；路由组件采用动态 import 实现懒加载，首屏仅加载首页所需资源；列表数据采用分类切换而非一次性全量渲染。以上措施可有效控制页面体积与渲染开销，性能约束可行。')

add_heading('2.4.3 合规伦理约束', level=3)
add_para('本项目为教学演示项目，所有 JSON 内课程、题库、示例代码、问答知识库均为教学演示内容或公开教学素材，不包含侵权资源、盗版内容或敏感数据；示例数据使用虚构或公开内容，符合学术诚信要求。AI 问答模块的 DeepSeek API Key 仅保存在用户浏览器 localStorage 中，不上传服务器或写入代码仓库，未配置 Key 时库外问题返回降级提示，不涉及用户隐私采集。')

add_heading('2.4.4 AI 工具使用边界', level=3)
add_para('AI 开发工具仅用于代码辅助生成、需求分析与问题排查；所有 AI 生成代码经本人理解、审查与验证后使用，并在技术复盘报告中如实记录使用情况。AI 问答模块对接的 DeepSeek API 仅用于回答编程学习相关问题，不用于生成违规内容。')

add_heading('2.5 风险分析与应对', level=2)
add_table(
    ['风险', '影响', '应对措施'],
    [
        ['SPA 部署后刷新 404', '深层路由无法直接访问', '构建后自动生成 404.html（复制 index.html），利用 GitHub Pages 404 回退机制实现 SPA 路由回退'],
        ['JSON 数据量大导致卡顿', '页面加载慢、交互不流畅', '路由懒加载、分类切换渲染、Vite 代码分割与压缩'],
        ['AI 接口异常或 Key 未配置', 'AI 问答功能不可用', '优先本地知识库检索，API 失败时自动降级返回提示，Key 可在页面随时配置'],
        ['路由守卫绕过', '未登录用户访问受保护页', '全局 beforeEach 守卫校验 localStorage token，未登录重定向到登录页并携带 redirect 参数'],
        ['本地存储数据丢失', '用户学习记录丢失', '关键操作即时写入 localStorage，提供个人中心查看与管理入口'],
    ]
)

doc.add_page_break()

# ========== 第3章 ==========
add_heading('3、系统架构设计', level=1)

add_heading('3.1 系统总体架构', level=2)
add_para('系统采用前端分层架构，自顶向下分为四层：')
add_para('（1）页面层（views）：负责路由与页面组织，包含 9 个视图组件——LoginView、RegisterView、ForgotPasswordView、HomeView、CoursesView、QuizView、ExamplesView、AssistantView、ProfileView，每个视图对应一个路由路径。', indent=True)
add_para('（2）组件层（components）：实现可复用 UI 组件，目前包含 NavBar 导航栏组件，根据路由 meta.hideNav 决定是否显示，在登录/注册/忘记密码页隐藏导航。', indent=True)
add_para('（3）状态与数据层（store + data + utils）：Pinia 的 user store 统一维护用户登录状态、本地用户库、验证码、学习方向等全局数据；src/data 目录下 4 个 JSON 文件提供课程、题库、示例代码、AI 知识库等业务数据；utils 目录提供校验（validate.js）、存储（storage.js）、AI 检索（assistant.js）等工具函数。', indent=True)
add_para('（4）工程与部署层：Vite 负责开发服务器与生产构建，vite.config.js 配置 GitHub Pages 部署 base 路径与代码分割；scripts/postbuild.mjs 在构建后生成 404.html 实现 SPA 回退；.github/workflows/deploy.yml 定义 GitHub Actions CI 流水线，自动构建并部署到 GitHub Pages。', indent=True)
add_para('分层架构使各模块职责单一、便于维护与扩展，页面层不直接操作底层数据，通过 store 与 utils 进行数据交互。')

add_heading('3.2 模块划分与职责', level=2)
add_table(
    ['层级', '目录/文件', '职责'],
    [
        ['入口', 'index.html, src/main.js', '应用入口，挂载 Pinia 与 Router，引入全局样式'],
        ['根组件', 'src/App.vue', '根组件，控制导航栏显示，渲染 router-view'],
        ['路由', 'src/router/index.js', '10 条路由定义，全局登录守卫，页面标题设置，滚动行为'],
        ['页面层', 'src/views/*.vue', '9 个业务页面视图组件'],
        ['组件层', 'src/components/NavBar.vue', '顶部导航栏，路由跳转，登录状态显示与退出'],
        ['状态层', 'src/store/user.js', 'Pinia store：用户注册/登录/登出/重置密码/学习方向管理'],
        ['数据层', 'src/data/*.json', 'courses/quiz/examples/assistant 四份业务数据'],
        ['工具层', 'src/utils/*.js', 'validate（手机号/表单校验）、storage（本地存储封装）、assistant（知识库检索与 AI 调用）'],
        ['样式', 'src/assets/main.css', '全局样式与 reset'],
        ['构建', 'vite.config.js, scripts/postbuild.mjs', 'Vite 配置、构建后 404.html 生成'],
        ['部署', '.github/workflows/deploy.yml', 'GitHub Actions CI/CD 流水线'],
    ]
)

add_heading('3.3 页面结构与路由', level=2)
add_table(
    ['路由路径', '页面名称', '组件', '是否需登录', '是否显示导航'],
    [
        ['/login', '登录', 'LoginView', '否', '否'],
        ['/register', '注册', 'RegisterView', '否', '否'],
        ['/forgot-password', '忘记密码', 'ForgotPasswordView', '否', '否'],
        ['/', '首页', 'HomeView', '否', '是'],
        ['/courses', '学习资源', 'CoursesView', '是', '是'],
        ['/quiz', '题库练习', 'QuizView', '是', '是'],
        ['/examples', '示例代码', 'ExamplesView', '是', '是'],
        ['/assistant', 'AI 问答', 'AssistantView', '是', '是'],
        ['/profile', '个人中心', 'ProfileView', '是', '是'],
        ['/:pathMatch(.*)*', '404 重定向', '—', '—', '—'],
    ]
)
add_para('路由采用 createWebHistory 模式，全局前置守卫（beforeEach）检查 localStorage 中的 webcreate_token，对 meta.requiresAuth 为 true 的路由进行登录校验，未登录时重定向到 /login 并携带 redirect 参数以便登录后回跳。全局后置钩子（afterEach）根据路由 meta.title 设置浏览器页面标题。')

add_heading('3.4 关键业务流程', level=2)

add_heading('3.4.1 用户注册与登录流程', level=3)
add_para('注册流程：用户输入手机号 → 前端校验手机号格式（isValidPhone）→ 检查本地用户库中手机号是否已注册 → 点击获取验证码 → 生成 6 位验证码并启动 60 秒倒计时 → 用户输入验证码 → 校验验证码正确性与有效期 → 设置密码 → 写入本地用户库 → 注册成功跳转登录页。')
add_para('登录流程：用户输入手机号与密码 → 前端校验手机号格式 → 查询本地用户库 → 校验密码 → 生成登录 token 写入 localStorage → 跳转首页或 redirect 指定页面。登录态通过路由守卫全局校验。')

add_heading('3.4.2 AI 问答检索流程', level=3)
add_para('用户输入问题 → 首先在本地知识库（assistant.json，45 条）中进行关键词匹配检索 → 检索算法采用分词匹配 + 权重打分（强词权重 > 弱词权重 > 先出现优先级），仅当问题包含"本项目/本网站/WebCreate"等限定词时才命中项目相关条目，避免普适问题误抢答 → 若本地检索命中且置信度足够，直接返回本地答案 → 若未命中或用户需要 AI 回答，检查 localStorage 中是否配置了 DeepSeek API Key → 已配置则调用 DeepSeek API 获取回答 → 未配置或 API 调用失败则返回降级提示，建议用户配置 Key 或参考本地知识库。对话记录保存在组件状态中，清空对话后欢迎卡片重现。')

add_heading('3.4.3 题库练习流程', level=3)
add_para('用户选择知识点 → 从 quiz.json 中筛选对应知识点的题目 → 随机抽题展示 → 用户作答 → 即时判题并显示正确答案与解析 → 答错的题目自动加入错题本（写入 localStorage）→ 支持题目收藏 → 支持抽题循环玩法，答完当前题目后自动加载下一题 → 个人中心可查看错题本与学习记录。')

doc.add_page_break()

# ========== 第4章 ==========
add_heading('4、环境搭建与工具使用', level=1)

add_heading('4.1 运行环境', level=2)
add_table(
    ['类别', '环境/工具', '版本'],
    [
        ['操作系统', 'Windows', 'Windows 10/11'],
        ['运行时', 'Node.js', '20 LTS 及以上'],
        ['包管理器', 'npm', '随 Node.js 安装'],
        ['前端框架', 'Vue', '3.4.29'],
        ['构建工具', 'Vite', '5.3.1'],
        ['路由', 'Vue Router', '4.3.3'],
        ['状态管理', 'Pinia', '2.1.7'],
        ['开发工具', 'VS Code', '最新版（Volar / ESLint / Prettier）'],
        ['版本管理', 'Git', '最新版'],
        ['部署平台', 'GitHub Pages', '—'],
        ['CI/CD', 'GitHub Actions', 'ubuntu-latest'],
        ['浏览器', 'Chrome / Edge', '最新版'],
    ]
)

add_heading('4.2 安装与配置步骤', level=2)
add_para('（1）安装 Node.js：从官网下载 LTS 版本（20.x），安装时勾选 "Add to PATH"，安装完成后在终端执行 node -v 与 npm -v 验证。', indent=True)
add_para('（2）初始化项目：使用 npm create vite@latest 命令创建 Vite + Vue 项目，项目名 webcreate，选择 Vue 模板与 JavaScript 变体。', indent=True)
add_para('（3）安装依赖：进入项目目录执行 npm install 安装基础依赖；执行 npm install vue-router@4 pinia 安装路由与状态管理库。', indent=True)
add_para('（4）配置 Vite：编辑 vite.config.js，配置 @ 别名指向 src 目录，配置 build 时 base 为 /webcreate/（GitHub Pages 项目站点路径），配置 manualChunks 将 vue 核心库独立打包。', indent=True)
add_para('（5）配置 Git：执行 git init 初始化仓库，创建 .gitignore 排除 node_modules、dist 等目录，关联 GitHub 远程仓库。', indent=True)
add_para('（6）创建三分支：基于 main 创建 develop 与 test 分支，日常开发在 develop 分支进行，测试部署合并到 test，正式发布合并到 main。', indent=True)

add_heading('4.3 启动方式', level=2)
add_para('本地开发模式：执行 npm run dev，Vite 启动开发服务器（默认端口 5173），浏览器访问 http://localhost:5173 即可，支持热更新。', indent=False)
add_para('生产构建：执行 npm run build，Vite 构建生产产物到 dist/ 目录，构建完成后自动执行 postbuild.mjs 生成 404.html。', indent=False)
add_para('本地预览构建结果：执行 npm run preview，该命令先以 base / 重新构建（build:local）再启动 vite preview，解决直接预览部署版资源路径 404 导致的空白页问题。', indent=False)

add_heading('4.4 配置中遇到的问题与解决', level=2)
add_table(
    ['问题', '原因', '解决方法'],
    [
        ['vite preview 预览空白页', 'npm run build 以 base /webcreate/ 构建，vite preview 直接预览时资源路径带 /webcreate/ 前缀导致 404', '新增 build:local 脚本以 base / 构建，preview 脚本先执行 build:local 再启动预览（git 提交 9c240d9）'],
        ['GitHub Pages 深层路由刷新 404', 'GitHub Pages 不支持 SPA rewrites，history 模式深层路由刷新返回 404', '新增 scripts/postbuild.mjs，构建后将 index.html 复制为 404.html，利用 Pages 404 回退机制（git 提交 d831e6b）'],
        ['GitHub Pages 部署后资源 404', '构建产物资源路径未配置 base，项目站点路径为 /webcreate/', 'vite.config.js 中配置 build 时 base 为 /webcreate/，dev 时保持 /（git 提交 d831e6b）'],
        ['develop 分支推送触发不必要部署', '初始 workflow 监听所有分支', '修改 workflow 仅监听 test 和 main 分支，develop 仅用于开发（git 提交 c0fe3dd）'],
    ]
)

add_heading('4.5 环境验证', level=2)
add_para('运行 npm run dev 打开默认页面，可正常浏览首页、登录/注册、各功能模块页面，即视为开发环境配置成功；执行 npm run build 能成功产出 dist 目录且包含 index.html 与 404.html，即视为构建环境配置成功；将代码推送到 test 分支后 GitHub Actions 自动构建部署，线上地址 https://xyy7539518246-stack.github.io/webcreate/ 可正常访问且深层路由刷新不 404，即视为部署环境配置成功。')

doc.add_page_break()

# ========== 第5章 ==========
add_heading('5、开发过程与问题解决', level=1)

add_heading('5.1 开发时间线', level=2)
add_para('项目开发周期为 2026 年 9 月 4 日至 9 月 6 日，共 3 天，全部在 develop 分支进行功能开发，完成后合并到 test 分支进行测试部署。以下为基于 Git 提交记录的开发时间线：')

add_heading('5.1.1 第一天（9月4日）：项目初始化与骨架搭建', level=3)
add_para('16:45 提交 4cc2a27「feat: 初始化项目」——使用 Vite 创建 Vue 3 项目，配置基础依赖。')
add_para('17:03 提交 91188ec「feat: 搭建项目主体结构与基础页面骨架」——创建 src 目录结构（views / components / store / router / utils / data / assets），配置 Vue Router 与 Pinia，搭建 App.vue 根组件与 NavBar 导航栏，创建各页面基础骨架。')

add_heading('5.1.2 第二天（9月5日）：核心业务模块开发', level=3)
add_para('上午至晚间集中开发用户认证、学习资源、题库练习三大模块：')
add_para('21:37 提交 ec15f35「feat: 新增手机号格式校验工具 isValidPhone」——在 utils/validate.js 中实现手机号正则校验。')
add_para('21:37 提交 da2121b「feat: 重构用户 store，支持本地用户库、验证码与登录凭证」——重写 store/user.js，实现用户注册、登录、登出、验证码生成与校验、学习方向管理等 Pinia actions。')
add_para('21:37 提交 392f98a「feat: 新增注册页，支持手机号唯一性校验、6位验证码与60s倒计时」——完成 RegisterView.vue，包含手机号格式校验、本地用户库查重、验证码生成与倒计时、密码设置。')
add_para('21:37 提交 43675e2「feat: 登录页改为手机号登录，未注册时提示并引导跳转注册页」——改造 LoginView.vue，支持手机号登录、未注册引导、登录后回跳。')
add_para('21:44 提交 970742d「feat: 新增学习资源静态数据 courses.json」——创建课程数据文件，包含知识点分类、学习路线、课程与竞赛内容。')
add_para('21:44 提交 c4420d0「feat: 学习资源页改为列表渲染与分类切换」——完成 CoursesView.vue，实现路线图时间线、课程/竞赛过滤 tab。')
add_para('21:50 提交 3067fc3「feat: 用户 store 新增学习方向字段与 setDirection 持久化方法」——扩展用户状态。')
add_para('21:50 提交 fd611af「feat: 首页支持选择学习方向并展示对应路线，个人中心可查看修改」——完成 HomeView.vue 学习方向选择功能。')
add_para('21:54 提交 d61aa18「feat: 学习资源扩展为五类资源类型，新增 Java 后端与前端进阶两条路线」——丰富 courses.json 数据。')
add_para('21:54 提交 ca74f77「feat: 课程/竞赛页类型过滤 tab 改为按数据动态生成」——优化过滤逻辑，适配新资源类型。')
add_para('21:57 提交 f237440「fix: 学习方向校验白名单改为动态读取 courses.json」——修复新增路线无法选择的问题。')
add_para('22:56 提交 28f9f9a「feat: 题库抽题循环玩法、题目收藏与个人中心学习记录页」——完成 QuizView.vue 与 ProfileView.vue 核心功能。')
add_para('23:08 提交 fc0c70e「feat: store 新增 resetPassword 动作」——支持验证码通过后重置密码。')
add_para('23:09 提交 04bb3e6「feat: 新增忘记密码页」——完成 ForgotPasswordView.vue。')
add_para('23:09 提交 986681a「feat: 登录页新增忘记密码入口链接」。')
add_para('23:27 合并 PR #1，将 develop 合并到 test 进行首次测试部署。')
add_para('23:29 起连续完成部署配置：d831e6b 配置 GitHub Pages base 与 404.html 回退、adb0195 新增 GitHub Actions workflow、26b0528 更新部署说明、575a8d3 补充部署地址、c0fe3dd 修改 workflow 监听分支。')
add_para('23:58 提交 0917377「chore: re-trigger Pages deployment after environment config」——配置 GitHub Pages 环境后重新触发部署。')

add_heading('5.1.3 第三天（9月6日）：示例代码、AI 问答与响应式适配', level=3)
add_para('18:19 提交 70cdaa2「feat: 完成示例代码模块」——创建 examples.json，完成 ExamplesView.vue，支持分类切换、语法高亮、一键复制。')
add_para('18:59 提交 a8b3363「feat: 完成 AI 问答模块」——创建 assistant.json 与 utils/assistant.js，完成 AssistantView.vue，实现本地知识库检索 + DeepSeek 直连 + 降级提示与对话记录。')
add_para('18:59 提交 9c240d9「fix: 修复 vite preview 预览空白页」——新增 build:local 脚本。')
add_para('19:43 提交 b8565fd「feat: AI 问答支持页面配置 API Key」——移除硬编码 .env，Key 存 localStorage，彻底避免 key 泄露。')
add_para('19:51 提交 2095bbf「feat: 丰富 AI 问答本地知识库（20→45 条）」——新增 CSS/JS 进阶、Vue3 进阶、C++ 算法数据结构、Java、Git、蓝桥杯、复杂度分析等条目。')
add_para('19:59 提交 c009612「fix: 修正 AI 问答检索逻辑」——普适问题不抢答本项目条目，去掉 question 前 6 字误加分，新增同分优先级（强词>弱词>先出现），修复技术栈误中栈队列、项目目录结构误中等边界问题。')
add_para('20:04 提交 716e65e「feat: AI 问答欢迎页」——对话为空时显示问候语+推荐问题快捷入口，清空对话后欢迎卡片重现。')
add_para('20:20 提交 6b3a678「fix: 注册/忘记密码页验证码改为页面内显示，去掉 alert 阻塞」——原 alert 阻塞导致验证码已开始过期但倒计时未启动，正确验证码偶发报错；改为页面内显示后倒计时与有效期完全同步。')
add_para('20:21 合并 PR #2。')
add_para('21:44 提交 54544a6「feat: 全站响应式适配，支持手机与电脑端访问」——通过媒体查询调整全站布局，完成移动端适配。')
add_para('21:44 合并 develop 到 test，完成测试部署。')

add_heading('5.2 典型问题与解决过程', level=2)

add_heading('5.2.1 验证码倒计时与有效期不同步问题', level=3)
add_para('问题描述：注册页与忘记密码页最初使用 alert 弹窗显示验证码，alert 会阻塞 JavaScript 执行线程，导致验证码已生成并开始过期，但 60 秒倒计时在用户关闭 alert 后才启动，造成倒计时与验证码实际有效期不同步，用户在倒计时显示未结束时输入正确验证码却偶发报错。')
add_para('解决过程：定位到 alert 阻塞是根本原因后，将验证码显示方式改为页面内文本展示（git 提交 6b3a678），去掉 alert 阻塞，使验证码生成、倒计时启动、有效期计算完全同步在同一事件循环中。修改后验证码倒计时与有效期严格一致，用户体验与正确性均得到保障。')
add_para('经验总结：在涉及计时器与异步状态的场景中，应避免使用阻塞式 UI（alert/confirm/prompt），改用非阻塞的页面内展示方式，确保状态时序的一致性。')

add_heading('5.2.2 AI 问答检索误匹配问题', level=3)
add_para('问题描述：AI 问答本地知识库检索最初存在两个问题——一是普适性问题（如"什么是栈"）会误抢答本项目相关条目（如"本项目技术栈"），因为关键词"栈"同时匹配；二是检索算法对 question 前 6 个字额外加分，导致短问题评分异常。此外还存在技术栈误中"栈/队列"、项目目录结构误中等边界问题。')
add_para('解决过程：（git 提交 c009612）重构检索逻辑——增加限定词判断，仅当问题包含"本项目/本网站/WebCreate"等限定词时才命中项目相关条目，普适问题走通用知识匹配；去掉 question 前 6 字误加分逻辑；新增同分优先级规则（强词权重 > 弱词权重 > 先出现优先级），确保匹配结果的合理性。修复后检索准确率显著提升。')
add_para('经验总结：关键词匹配检索需要精心设计权重与边界规则，简单的字符串包含匹配会导致大量误报，需要结合语义限定词、权重分级与优先级规则来提升检索质量。')

add_heading('5.2.3 vite preview 预览空白页问题', level=3)
add_para('问题描述：执行 npm run build 后用 vite preview 预览构建结果，页面显示空白。原因是 vite.config.js 中配置了 build 时 base 为 /webcreate/（用于 GitHub Pages 项目站点），构建产物中所有资源路径都带有 /webcreate/ 前缀，而 vite preview 默认以根路径提供服务，导致资源请求 404。')
add_para('解决过程：（git 提交 9c240d9）新增 build:local 脚本（vite build --base=/），并将 preview 脚本改为 npm run build:local && vite preview，即以根路径重新构建后再启动预览。这样本地预览不受 GitHub Pages base 配置影响，线上部署仍使用 npm run build（base /webcreate/），两者互不干扰。')
add_para('经验总结：当项目需要同时支持本地预览与子路径部署时，应通过不同的构建脚本区分 base 配置，避免单一配置导致某一环境下资源路径错误。')

add_heading('5.2.4 学习方向白名单硬编码问题', level=3)
add_para('问题描述：首页学习方向选择功能最初将可选方向硬编码在组件中，当 courses.json 新增学习路线（如 Java 后端、前端进阶）后，首页选择器无法显示新增方向，导致用户无法选择新课程路线。')
add_para('解决过程：（git 提交 f237440）将学习方向校验白名单改为动态读取 courses.json 中的路线数据，组件初始化时从数据文件提取所有可用方向，确保数据与 UI 保持同步。后续新增路线只需修改 JSON 数据，无需改动组件代码。')
add_para('经验总结：可配置项应从数据源动态读取，避免在代码中硬编码与数据重复的常量，否则数据变更时需要同步修改多处代码，容易遗漏。')

doc.add_page_break()

# ========== 第6章 ==========
add_heading('6、部署与功能测试', level=1)

add_heading('6.1 部署过程', level=2)
add_para('项目采用 GitHub Actions 自动构建部署到 GitHub Pages，部署流程如下：')
add_para('（1）代码推送到 test 或 main 分支，触发 .github/workflows/deploy.yml 中定义的 CI 流水线。', indent=True)
add_para('（2）CI 运行在 ubuntu-latest 环境，执行 actions/checkout@v4 检出代码，actions/setup-node@v4 配置 Node.js 20 环境并启用 npm 缓存。', indent=True)
add_para('（3）执行 npm ci 安装依赖（使用 package-lock.json 确保依赖版本一致）。', indent=True)
add_para('（4）执行 npm run build 构建生产产物到 dist/ 目录，构建完成后自动执行 postbuild.mjs 生成 404.html。', indent=True)
add_para('（5）使用 actions/upload-pages-artifact@v3 上传 dist 目录为 Pages 构件。', indent=True)
add_para('（6）deploy 作业使用 actions/deploy-pages@v4 将构件部署到 GitHub Pages。', indent=True)
add_para('（7）部署完成后访问线上地址：https://xyy7539518246-stack.github.io/webcreate/')
add_para('流水线配置了 concurrency 组（pages），同一时间只允许一次部署，新的部署会取消进行中的部署，避免并发部署冲突。GITHUB_TOKEN 授予 contents:read、pages:write、id-token:write 权限，满足 Pages 部署所需的最小权限。')
add_para('首次部署前需在仓库 Settings → Pages 中配置 Build and deployment → Source 为 GitHub Actions，并在 Settings → Environments → github-pages 中将 Deployment branches 配置为允许 test 和 main 分支。')

add_heading('6.2 功能测试', level=2)
add_para('项目开发完成后，对各功能模块进行了系统性功能测试，测试覆盖以下功能点：')

add_heading('6.2.1 用户认证模块测试', level=3)
add_para('测试内容包括：手机号格式校验（非法手机号拦截）、注册时手机号唯一性校验（已注册手机号拦截）、验证码生成与 60 秒倒计时、验证码正确性校验、验证码过期校验、密码设置与确认、登录成功与失败（密码错误、未注册）、登录后 token 写入 localStorage、路由守卫（未登录访问受保护页面重定向到登录页）、登录后回跳 redirect 参数、退出登录清除 token、忘记密码验证码重置密码流程。')

add_heading('6.2.2 首页与学习资源模块测试', level=3)
add_para('测试内容包括：首页学习方向选择与持久化（刷新后保持选择）、首页功能导航跳转、学习资源页路线图时间线渲染、课程/竞赛类型过滤 tab 切换、新增资源类型后 tab 动态生成、学习方向白名单与 courses.json 数据同步。')

add_heading('6.2.3 题库练习模块测试', level=3)
add_para('测试内容包括：按知识点筛选题目、随机抽题、题目展示（题干/选项/解析）、即时判题（正确/错误反馈）、错题自动收集到 localStorage、错题本查看、题目收藏与取消收藏、抽题循环玩法（答完自动下一题）、个人中心学习记录展示。')

add_heading('6.2.4 示例代码模块测试', level=3)
add_para('测试内容包括：C++/Java/前端分类切换、代码语法高亮显示、一键复制功能（复制到剪贴板）、代码内容与 examples.json 数据一致性。')

add_heading('6.2.5 AI 问答模块测试', level=3)
add_para('测试内容包括：本地知识库检索（45 条条目命中测试）、普适问题不误抢答项目条目、限定词（本项目/WebCreate）命中项目条目、同分优先级规则、API Key 页面配置与 localStorage 保存、未配置 Key 时库外问题降级提示、DeepSeek API 调用成功与失败降级、对话记录展示、清空对话后欢迎卡片重现、推荐问题快捷入口点击命中本地库。')

add_heading('6.2.6 响应式与兼容性测试', level=3)
add_para('测试内容包括：Chrome 浏览器桌面端布局与交互、Edge 浏览器桌面端布局与交互、移动端视口（≤768px）布局自适应、导航栏在移动端的显示、各页面在窄屏下的排版与可读性、触控操作友好性。')

add_heading('6.2.7 部署与路由测试', level=3)
add_para('测试内容包括：GitHub Actions 自动构建成功、线上地址可正常访问、首页加载无空白、各路由页面可正常跳转、深层路由刷新不 404（404.html 回退生效）、构建产物资源路径正确（base /webcreate/）、test 与 main 分支部署覆盖逻辑。')

add_heading('6.3 测试结果', level=2)
add_para('经过上述系统性测试，各功能模块核心功能均正常运行。测试过程中发现并修复了以下问题：')
add_para('（1）验证码 alert 阻塞导致倒计时与有效期不同步——已修复（改为页面内显示）。', indent=True)
add_para('（2）AI 问答普适问题误抢答项目条目——已修复（增加限定词判断与权重规则）。', indent=True)
add_para('（3）vite preview 预览空白页——已修复（新增 build:local 脚本）。', indent=True)
add_para('（4）学习方向白名单硬编码导致新增路线无法选择——已修复（动态读取 courses.json）。', indent=True)
add_para('（5）GitHub Pages 深层路由刷新 404——已修复（postbuild 生成 404.html）。', indent=True)
add_para('修复后回归测试全部通过，线上部署环境运行稳定，各页面可正常访问与交互。')

doc.add_page_break()

# ========== 第7章 ==========
add_heading('7、工具/框架局限性分析', level=1)

add_heading('7.1 Vue 3 框架局限性', level=2)
add_para('Vue 3 的组合式 API（Composition API）虽然提供了更好的逻辑复用与类型推断能力，但也带来了一定的学习成本，初学者需要理解 ref/reactive、computed、watch、生命周期钩子等概念，与选项式 API（Options API）的思维方式不同。对于本项目这种中小型课程实践项目，部分简单页面使用组合式 API 可能略显繁琐，原生 JavaScript 或更轻量的方案可能更简洁。')
add_para('此外，Vue 作为前端框架，其响应式系统在处理大量数据（如题库数百道题目一次性渲染）时可能存在性能瓶颈，需要通过分页、虚拟滚动等方式优化。本项目通过分类切换与按需渲染缓解了这一问题，但未实现虚拟滚动。')
add_para('Vue 的单文件组件（SFC）需要构建工具支持，不能直接在浏览器中运行，这增加了构建与部署的复杂度。虽然 Vite 极大地提升了开发体验，但对于不需要构建的简单场景，原生 HTML/CSS/JS 可能更直接。')

add_heading('7.2 Vite 构建工具局限性', level=2)
add_para('Vite 基于原生 ES 模块（ESM）实现开发服务器，在开发模式下速度极快，但在生产构建时仍使用 Rollup 进行打包，对于超大型项目的构建时间可能较长。本项目规模较小，构建时间在可接受范围内。')
add_para('Vite 的配置虽然灵活，但对于初学者来说，base 路径、代理、代码分割等配置需要一定的学习成本。本项目在部署过程中就遇到了 base 路径配置导致的预览空白页问题，需要通过多脚本方案解决。')
add_para('此外，Vite 对 CommonJS 模块的兼容依赖于预构建，某些老旧的第三方库可能在 Vite 环境下出现兼容问题。本项目仅使用了 Vue 生态的官方库，未遇到此类问题。')

add_heading('7.3 静态托管（GitHub Pages）局限性', level=2)
add_para('GitHub Pages 作为纯静态托管平台，仅支持前端静态资源（HTML/CSS/JS/图片），无法承载后端服务、数据库或服务端渲染。这意味着本项目的用户数据只能存储在浏览器 localStorage 中，无法实现多设备同步、多用户共享数据、服务端校验等功能。')
add_para('GitHub Pages 不支持 SPA rewrites（即无法配置服务端路由回退），本项目通过构建后生成 404.html 的 hack 方式实现了 SPA 路由回退，但这种方式在某些边缘场景（如自定义 404 页面与应用路由冲突）下可能存在问题。')
add_para('GitHub Pages 有仓库大小与流量限制（建议仓库不超过 1GB，月度带宽不超过 100GB），对于大流量或大资源的项目不适用。本项目为教学演示项目，资源体积小、流量低，不受影响。')
add_para('GitHub Pages 仅支持 HTTPS，不支持自定义后端 API 代理，如果前端需要调用不支持 CORS 的第三方 API，需要额外的代理服务。本项目的 DeepSeek API 支持跨域，可直接从浏览器调用，未受此限制。')

add_heading('7.4 纯前端方案局限性', level=2)
add_para('本项目采用纯前端实现，所有用户数据（账号、密码、学习记录、错题、收藏）均存储在浏览器 localStorage 中，存在以下局限：')
add_para('（1）数据安全性：localStorage 中的数据可被用户在浏览器开发者工具中直接查看和修改，密码虽然经过简单处理但并非安全加密，不适合存储敏感信息。本项目为教学演示，不涉及真实用户数据，但在生产环境中必须使用后端服务进行身份认证与数据加密存储。', indent=True)
add_para('（2）数据持久性：localStorage 数据与浏览器绑定，清除浏览器数据、更换设备或使用隐私模式都会导致数据丢失，无法实现多设备同步。', indent=True)
add_para('（3）存储容量：localStorage 通常有 5-10MB 的容量限制，大量题库或学习记录可能超出容量。本项目数据量较小，未触及限制。', indent=True)
add_para('（4）无法实现服务端逻辑：如邮件验证码发送、真实短信验证、数据统计分析、多用户排行榜等需要服务端支持的功能无法实现。本项目的验证码为前端本地生成，仅作演示。', indent=True)

add_heading('7.5 AI 开发工具局限性', level=2)
add_para('AI 开发工具（如代码辅助生成、问题排查工具）在本项目开发中发挥了重要作用，但也存在以下局限：')
add_para('（1）生成代码可能存在逻辑错误或版本不匹配：AI 生成的代码可能基于旧版本 API 或存在边界条件遗漏，需要人工审查与验证。本项目中 AI 生成的检索算法初始版本存在误匹配问题，经过人工调试与重构后才达到预期效果。', indent=True)
add_para('（2）AI 对项目上下文的理解有限：AI 工具难以完全理解项目的整体架构与隐含约束，生成的代码可能与现有代码风格或架构不一致，需要人工调整。', indent=True)
add_para('（3）AI 无法替代调试与测试：AI 可以提供问题排查思路，但实际的复现、定位、修复与回归测试仍需开发者亲自完成。本项目中的验证码倒计时同步问题、vite preview 空白页问题等，都是通过实际运行与调试才定位到根本原因。', indent=True)
add_para('（4）AI 生成内容需要遵守使用边界：AI 工具仅用于辅助开发，所有生成内容必须经过本人理解、审查与验证后才能使用，不能直接照搬。本项目在技术复盘中如实记录了 AI 工具的使用情况。')

doc.add_page_break()

# ========== 第8章 ==========
add_heading('8、总结与改进方向', level=1)

add_heading('8.1 实践总结', level=2)
add_para('本次《Web开发实践》课程实践，我独立完成了"WebCreate —— AI 编程学习助手"前端单页应用的全流程开发，从需求分析、架构设计、编码实现、问题调试到线上部署与功能测试，历时 3 天，共完成 30 余次 Git 提交，实现了 7 大功能模块、9 个页面视图、4 份业务数据文件、完整的 CI/CD 部署流水线。')
add_para('通过本次实践，我在以下方面获得了显著提升：')
add_para('（1）Vue 3 组件化开发：深入理解了组合式 API（Composition API）的设计理念，掌握了 ref/reactive、computed、watch、生命周期钩子等核心概念，能够熟练拆分组件与复用逻辑。', indent=True)
add_para('（2）SPA 路由与状态管理：掌握了 Vue Router 的路由配置、动态路由、路由守卫（beforeEach/afterEach）、滚动行为定制等高级特性；掌握了 Pinia 状态管理的 store 定义、actions、getters 与持久化方案。', indent=True)
add_para('（3）前端工程化：掌握了 Vite 的配置与优化，包括路径别名、base 配置、代码分割（manualChunks）、构建后处理脚本（postbuild）等；理解了开发环境与生产环境的差异与配置分离。', indent=True)
add_para('（4）CI/CD 与静态部署：掌握了 GitHub Actions 工作流的编写，包括触发条件、权限配置、并发控制、Node 环境配置、npm ci 安装、构建与部署步骤；掌握了 GitHub Pages 的部署配置与 SPA 路由回退方案。', indent=True)
add_para('（5）Git 版本管理与分支策略：实践了 develop—test—main 三分支工作流，掌握了功能分支开发、Pull Request 合并、提交信息规范（feat/fix/docs/ci/chore）等工程化实践。', indent=True)
add_para('（6）问题排查与调试能力：在开发过程中遇到并解决了验证码倒计时同步、AI 检索误匹配、vite preview 空白页、学习方向白名单硬编码、SPA 路由 404 等多个典型问题，提升了问题定位与调试能力。', indent=True)
add_para('（7）复杂工程约束考量：在开发过程中系统考虑了跨浏览器兼容与响应式、性能优化、合规伦理、AI 工具使用边界等复杂工程约束，并落实到具体的技术措施中。', indent=True)

add_heading('8.2 改进方向', level=2)
add_para('基于本次实践的经验与局限性分析，后续可从以下方向进行改进与扩展：')
add_para('（1）接入后端服务与数据库：将用户认证、学习记录、错题本、收藏等数据迁移到后端服务（如 Node.js + Express + MongoDB/MySQL），实现多用户数据持久化与多设备同步，提升数据安全性。可采用前后端分离架构，前端通过 RESTful API 与后端交互。', indent=True)
add_para('（2）采用 UI 组件库提升界面质量：引入 Element Plus 或 Ant Design Vue 等成熟组件库，统一按钮、表单、弹窗、表格等组件的样式与交互，提升界面的专业度与一致性，减少自定义样式的维护成本。', indent=True)
add_para('（3）优化首屏加载与按需引入：进一步分析构建产物体积，对大型依赖（如代码高亮库）进行按需引入；实现路由级别的代码分割与预加载；对图片等静态资源进行压缩与懒加载；考虑使用 Service Worker 实现 PWA 离线访问。', indent=True)
add_para('（4）丰富题库与学习资源：持续扩充 quiz.json 题库数据，增加题目类型（多选、判断、编程题），实现题目难度分级与知识点关联分析；扩充 courses.json 学习路线，增加更多技术方向与阶段目标。', indent=True)
add_para('（5）增强 AI 问答能力：优化本地知识库检索算法，引入向量检索或更精细的语义匹配；增加对话历史上下文管理，支持多轮对话；增加 Markdown 渲染与代码高亮，提升 AI 回答的可读性；增加用户反馈机制（有用/无用），持续优化回答质量。', indent=True)
add_para('（6）增加数据可视化与学习分析：在个人中心增加学习数据统计图表（如答题正确率趋势、知识点掌握度雷达图、学习时长统计），使用 ECharts 等图表库实现数据可视化，帮助用户了解学习进度与薄弱环节。', indent=True)
add_para('（7）完善测试体系：引入 Vitest 进行单元测试，对工具函数（validate、storage、assistant 检索算法）编写测试用例；引入 Vue Test Utils 进行组件测试；考虑引入 E2E 测试（如 Playwright）进行关键流程的自动化测试，提升代码质量与回归效率。', indent=True)
add_para('（8）国际化与无障碍：增加 i18n 国际化支持，提供中英文切换；遵循 WCAG 无障碍标准，增加键盘导航、屏幕阅读器支持、颜色对比度优化，提升应用的可访问性。', indent=True)

doc.add_page_break()

# ========== 参考文献 ==========
add_heading('参考文献', level=1)
refs = [
    '[1] 尤雨溪. Vue.js 3 官方文档[EB/OL]. https://vuejs.org/, 2024.',
    '[2] MDN Web Docs. Web 开发技术文档[EB/OL]. https://developer.mozilla.org/zh-CN/, 2024.',
    '[3] Vite 官方文档. 下一代前端构建工具[EB/OL]. https://cn.vitejs.dev/, 2024.',
    '[4] Pinia 官方文档. Vue 官方状态管理库[EB/OL]. https://pinia.vuejs.org/zh/, 2024.',
    '[5] Vue Router 官方文档. Vue.js 官方路由[EB/OL]. https://router.vuejs.org/zh/, 2024.',
    '[6] GitHub Docs. GitHub Pages 文档[EB/OL]. https://docs.github.com/zh/pages, 2024.',
    '[7] GitHub Docs. GitHub Actions 文档[EB/OL]. https://docs.github.com/zh/actions, 2024.',
    '[8] 蔡冰. Vue.js 3 开发详解[M]. 北京：清华大学出版社，2023.',
    '[9] DeepSeek API 文档. 深度求索开放平台[EB/OL]. https://platform.deepseek.com/, 2024.',
    '[10] 廖雪峰. Git 教程[EB/OL]. https://www.liaoxuefeng.com/wiki/896043488029600, 2024.',
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.hanging_indent = Cm(0.74)
    run = p.add_run(ref)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(10.5)

doc.add_page_break()

# ========== 附录 ==========
add_heading('附录', level=1)

add_heading('附录 A：Git 提交记录', level=2)
add_para('以下为项目开发过程中的完整 Git 提交记录（按时间倒序），共 30 余次功能提交，涵盖项目初始化、各模块开发、问题修复、部署配置等全过程：')

git_log = [
    ['54544a6', '2026-09-06 21:44', 'feat: 全站响应式适配，支持手机与电脑端访问'],
    ['6b3a678', '2026-09-06 20:20', 'fix: 注册/忘记密码页验证码改为页面内显示，去掉 alert 阻塞'],
    ['716e65e', '2026-09-06 20:04', 'feat: AI 问答欢迎页——问候语+推荐问题快捷入口'],
    ['c009612', '2026-09-06 19:59', 'fix: 修正 AI 问答检索逻辑——限定词判断+权重规则'],
    ['2095bbf', '2026-09-06 19:51', 'feat: 丰富 AI 问答本地知识库（20→45 条）'],
    ['b8565fd', '2026-09-06 19:43', 'feat: AI 问答支持页面配置 API Key（存 localStorage）'],
    ['9c240d9', '2026-09-06 18:59', 'fix: 修复 vite preview 预览空白页（build:local）'],
    ['a8b3363', '2026-09-06 18:59', 'feat: 完成 AI 问答模块（本地检索+DeepSeek+降级）'],
    ['70cdaa2', '2026-09-06 18:19', 'feat: 完成示例代码模块（分类切换/高亮/复制）'],
    ['0917377', '2026-09-05 23:58', 'chore: re-trigger Pages deployment after environment config'],
    ['c0fe3dd', '2026-09-05 23:45', 'ci: workflow 改为监听 test 和 main 分支'],
    ['575a8d3', '2026-09-05 23:38', 'docs: 补充 GitHub Pages 部署地址'],
    ['26b0528', '2026-09-05 23:29', 'docs: 更新 GitHub Pages 部署说明'],
    ['adb0195', '2026-09-05 23:29', 'ci: 新增 GitHub Actions 自动部署 workflow'],
    ['d831e6b', '2026-09-05 23:29', 'build: 配置 GitHub Pages 部署 base 与 404.html SPA 回退'],
    ['986681a', '2026-09-05 23:09', 'feat: 登录页新增忘记密码入口链接'],
    ['04bb3e6', '2026-09-05 23:09', 'feat: 新增忘记密码页，验证码重置密码'],
    ['fc0c70e', '2026-09-05 23:08', 'feat: store 新增 resetPassword 动作'],
    ['28f9f9a', '2026-09-05 22:56', 'feat: 题库抽题循环玩法、题目收藏与个人中心学习记录'],
    ['f237440', '2026-09-05 21:57', 'fix: 学习方向校验白名单改为动态读取 courses.json'],
    ['ca74f77', '2026-09-05 21:54', 'feat: 课程/竞赛页类型过滤 tab 改为按数据动态生成'],
    ['d61aa18', '2026-09-05 21:54', 'feat: 学习资源扩展为五类资源类型'],
    ['fd611af', '2026-09-05 21:50', 'feat: 首页支持选择学习方向并展示对应路线'],
    ['3067fc3', '2026-09-05 21:50', 'feat: 用户 store 新增学习方向字段与 setDirection'],
    ['c4420d0', '2026-09-05 21:44', 'feat: 学习资源页改为列表渲染与分类切换'],
    ['970742d', '2026-09-05 21:44', 'feat: 新增学习资源静态数据 courses.json'],
    ['43675e2', '2026-09-05 21:37', 'feat: 登录页改为手机号登录'],
    ['392f98a', '2026-09-05 21:37', 'feat: 新增注册页，手机号唯一性校验+验证码+倒计时'],
    ['da2121b', '2026-09-05 21:37', 'feat: 重构用户 store，支持本地用户库、验证码与登录凭证'],
    ['ec15f35', '2026-09-05 21:37', 'feat: 新增手机号格式校验工具 isValidPhone'],
    ['91188ec', '2026-09-04 17:03', 'feat: 搭建项目主体结构与基础页面骨架'],
    ['4cc2a27', '2026-09-04 16:45', 'feat: 初始化项目'],
]
add_table(['提交哈希', '时间', '提交说明'], git_log)

add_heading('附录 B：项目目录结构', level=2)
dir_structure = '''webcreate/
├── index.html                  # 入口 HTML
├── vite.config.js              # Vite 配置（base/别名/代码分割）
├── package.json                # 依赖与脚本
├── package-lock.json           # 依赖版本锁定
├── .gitignore                  # 忽略文件配置
├── README.md                   # 项目说明文档
├── .github/
│   └── workflows/
│       └── deploy.yml          # GitHub Actions CI/CD 流水线
├── scripts/
│   └── postbuild.mjs           # 构建后生成 404.html（SPA 回退）
├── public/                     # 静态资源（不经过构建）
├── src/
│   ├── main.js                 # 应用入口，挂载 Pinia 与 Router
│   ├── App.vue                 # 根组件，导航栏控制与路由视图
│   ├── assets/
│   │   └── main.css            # 全局样式与 reset
│   ├── components/
│   │   └── NavBar.vue          # 顶部导航栏组件
│   ├── views/                  # 页面层（9 个视图组件）
│   │   ├── LoginView.vue       # 登录页        /login
│   │   ├── RegisterView.vue    # 注册页        /register
│   │   ├── ForgotPasswordView.vue  # 忘记密码页 /forgot-password
│   │   ├── HomeView.vue        # 首页          /
│   │   ├── CoursesView.vue     # 学习资源页    /courses
│   │   ├── QuizView.vue        # 题库练习页    /quiz
│   │   ├── ExamplesView.vue    # 示例代码页    /examples
│   │   ├── AssistantView.vue   # AI 问答页     /assistant
│   │   └── ProfileView.vue     # 个人中心页    /profile
│   ├── router/
│   │   └── index.js            # 路由配置（10 路由+登录守卫+标题设置）
│   ├── store/
│   │   └── user.js             # Pinia 用户状态管理
│   ├── utils/
│   │   ├── validate.js         # 表单校验工具（手机号等）
│   │   ├── storage.js          # 本地存储封装
│   │   └── assistant.js        # AI 知识库检索与 API 调用
│   └── data/                   # 业务数据（JSON）
│       ├── courses.json        # 学习资源数据（18.8 KB）
│       ├── quiz.json           # 题库数据（46.2 KB）
│       ├── examples.json       # 示例代码数据（15.1 KB）
│       └── assistant.json      # AI 知识库数据（36.2 KB，45 条）
└── dist/                       # 构建产物（不纳入版本管理）'''

p = doc.add_paragraph()
p.paragraph_format.first_line_indent = Cm(0)
run = p.add_run(dir_structure)
run.font.name = 'Consolas'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
run.font.size = Pt(9)

# ========== 保存 ==========
output_path = r'D:\webcreat\webcreate\docs\WebCreate项目实践报告_v2.docx'
import os
os.makedirs(os.path.dirname(output_path), exist_ok=True)
doc.save(output_path)
print(f'报告已生成: {output_path}')
print(f'段落数: {len(doc.paragraphs)}')
