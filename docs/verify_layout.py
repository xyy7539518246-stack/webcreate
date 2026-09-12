# -*- coding: utf-8 -*-
import docx, os
from docx.shared import Pt, Cm

doc = docx.Document(r'D:\webcreat\webcreate\docs\WebCreate项目实践报告.docx')
ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

sec = doc.sections[0]
print('=== 页面设置 ===')
print(f'页面: {sec.page_width.cm:.1f} x {sec.page_height.cm:.1f} cm (A4=21.0x29.7)')
print(f'页边距: 上{sec.top_margin.cm:.2f} 下{sec.bottom_margin.cm:.2f} 左{sec.left_margin.cm:.2f} 右{sec.right_margin.cm:.2f} cm')

print('\n=== 标题样式 ===')
for lv in [1,2,3]:
    s = doc.styles[f'Heading {lv}']
    pf = s.paragraph_format
    b = pf.space_before.pt if pf.space_before else 0
    a = pf.space_after.pt if pf.space_after else 0
    print(f'Heading {lv}: {s.font.size.pt}pt {s.font.name} bold={s.font.bold} | before={b}pt after={a}pt align={pf.alignment}')

print('\n=== 正文段落抽样 ===')
for i in [50, 80, 120, 160, 200]:
    p = doc.paragraphs[i]
    pf = p.paragraph_format
    indent_pt = pf.first_line_indent.pt if pf.first_line_indent else 0
    print(f'[{i}] {p.style.name}: indent={indent_pt}pt line={pf.line_spacing} text="{p.text[:25]}"')

print('\n=== 表格样式 ===')
for ti, t in enumerate(doc.tables):
    hdr_cell = t.rows[0].cells[0]
    tcPr = hdr_cell._tc.tcPr
    shd = tcPr.find(ns+'shd') if tcPr is not None else None
    fill = shd.get(ns+'fill') if shd is not None else 'none'
    hdr_align = t.rows[0].cells[0].paragraphs[0].alignment
    runs = t.rows[0].cells[0].paragraphs[0].runs
    hdr_bold = runs[0].font.bold if runs else False
    print(f'表格{ti}: {len(t.rows)}行x{len(t.columns)}列 表头底色={fill} 对齐={hdr_align} 加粗={hdr_bold}')

f = os.path.getsize(r'D:\webcreat\webcreate\docs\WebCreate项目实践报告.docx')
print(f'\n文件大小: {f/1024:.1f} KB')
print(f'总段落: {len(doc.paragraphs)}  表格: {len(doc.tables)}')
