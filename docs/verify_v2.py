# -*- coding: utf-8 -*-
import docx
doc = docx.Document(r'D:\webcreat\webcreate\docs\WebCreate项目实践报告_v2.docx')

print('=== 封面区域（前20段）===')
for i in range(20):
    p = doc.paragraphs[i]
    text = p.text.strip()
    if not text:
        continue
    runs_info = ''
    if p.runs:
        r = p.runs[0]
        sz = r.font.size.pt if r.font.size else '?'
        nm = r.font.name
        bd = r.font.bold
        runs_info = f'[{nm} {sz}pt bold={bd}]'
    print(f'[{i}] align={p.alignment} {runs_info}')
    print(f'     {text[:60]}')

print('\n=== 摘要/目录标题 ===')
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t in ['摘　　要', '目　　录'] or '摘' in t and '要' in t and len(t) < 10:
        runs_info = ''
        if p.runs:
            r = p.runs[0]
            sz = r.font.size.pt if r.font.size else '?'
            nm = r.font.name
            bd = r.font.bold
            runs_info = f'[{nm} {sz}pt bold={bd}]'
        print(f'[{i}] style={p.style.name} align={p.alignment} {runs_info}')
        print(f'     文本: {t}')

print('\n=== 目录项抽样 ===')
for i in range(24, 36):
    p = doc.paragraphs[i]
    t = p.text.strip()
    if t:
        runs_info = ''
        if p.runs:
            r = p.runs[0]
            sz = r.font.size.pt if r.font.size else '?'
            nm = r.font.name
            runs_info = f'[{nm} {sz}pt]'
        print(f'[{i}] {runs_info} {t[:50]}')

import os
f = os.path.getsize(r'D:\webcreat\webcreate\docs\WebCreate项目实践报告_v2.docx')
print(f'\n文件大小: {f/1024:.1f} KB')
